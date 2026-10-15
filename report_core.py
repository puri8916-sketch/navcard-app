"""รายงานการปิดการ์ดนำร่อง — monthly workbook built from ReportCardIsClosedDaily.xlsx exports.

Each day ("tab") is stored as plain data: {date, sent, cards:[{no, agent, ship, imo, cs, loa, ev:[[when, draft, berth, pilot, closer], ...]}]}
and the whole month is rebuilt from those days with the look of the office's own workbook
(report_template.xlsx = the 13-10 tab of ต.ค.69 with its data cleared).
"""
import copy
import io
import os
import re
import zipfile
from datetime import date, timedelta

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Border, Side
from openpyxl.styles.colors import Color

_HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(_HERE, "report_template.xlsx")
TITLE = "รายงานการปิดการ์ดนำร่อง เขต เขตสงขลา"
SENT_RGB = "FF00B050"
TH_MONTHS = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]

# layout of report_template.xlsx (sheet TPL)
_FIRST, _MID, _LAST = 4, 7, 5       # card first row / middle continuation row / last continuation row
_SIG0, _SIG_N = 24, 6               # signature block = 6 rows starting right after the last data row
_SIG_MERGES = [(3, 2, 4), (4, 2, 6), (5, 2, 6), (4, 8, 10), (5, 8, 12), (6, 8, 12)]   # (row offset, first col, last col)
_SIG_IMG = dict(col=7, col_off=688363, row_off=160883, cx=1707300, cy=818512)        # anchored on the 1st row after the data


def _s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def month_file_name(y, m):
    return f"{TH_MONTHS[m - 1]}{(y + 543) % 100}.xlsx"


def tab_name(iso):
    d = date.fromisoformat(iso)
    return f"{d.day}-{d.month}"


# ---------------------------------------------------------------- reading
def _read_cards(ws, first_row, stop):
    cards = []
    for r in range(first_row, ws.max_row + 1):
        a = _s(ws.cell(r, 1).value).strip()
        h = _s(ws.cell(r, 8).value)
        if stop(ws, r):
            break
        if a:
            cards.append({"no": a, "agent": _s(ws.cell(r, 2).value), "ship": _s(ws.cell(r, 3).value),
                          "imo": _s(ws.cell(r, 5).value), "cs": _s(ws.cell(r, 6).value),
                          "loa": _s(ws.cell(r, 7).value), "ev": []})
        if h.strip() and cards:
            cards[-1]["ev"].append([h, _s(ws.cell(r, 9).value), _s(ws.cell(r, 10).value),
                                    _s(ws.cell(r, 13).value), _s(ws.cell(r, 14).value)])
    return cards


def _dmy(s):
    d, m, y = (int(x) for x in s.split("/"))
    return date(y, m, d).isoformat()


def parse_daily(data):
    """ReportCardIsClosedDaily.xlsx -> (from_iso, to_iso, cards)."""
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    head, title = None, ""
    for r in range(1, min(ws.max_row, 30) + 1):
        for c in range(1, 15):
            v = _s(ws.cell(r, c).value)
            if "ระหว่างวันที่" in v:
                title = v
            if v.strip() == "Card Number" and head is None:
                head = r
    if head is None:
        raise ValueError("ไม่พบหัวตาราง 'Card Number' — ไฟล์นี้ไม่ใช่ ReportCardIsClosedDaily")
    ds = re.findall(r"\d{1,2}/\d{1,2}/\d{4}", title)
    if not ds:
        raise ValueError("ไม่พบช่วงวันที่ (ระหว่างวันที่ … ถึงวันที่ …) ในไฟล์")
    d0, d1 = _dmy(ds[0]), _dmy(ds[-1])
    cards = _read_cards(ws, head + 1, lambda ws, r: False)
    if not cards:
        raise ValueError("ไม่พบรายการการ์ดในไฟล์")
    return d0, d1, cards


def _signature_from_zip(data):
    """First picture in the workbook, cropped the way Excel shows it (a:srcRect)."""
    try:
        from PIL import Image
        z = zipfile.ZipFile(io.BytesIO(data))
        for name in sorted(n for n in z.namelist() if re.match(r"xl/drawings/drawing\d+\.xml$", n)):
            x = z.read(name).decode("utf-8", "ignore")
            m = re.search(r'r:embed="(rId\d+)"', x)
            if not m:
                continue
            rels = z.read(name.replace("drawings/", "drawings/_rels/") + ".rels").decode("utf-8", "ignore")
            t = re.search(r'Id="%s"[^>]*Target="([^"]+)"' % m.group(1), rels) or \
                re.search(r'Target="([^"]+)"[^>]*Id="%s"' % m.group(1), rels)
            if not t:
                continue
            tgt = t.group(1)
            path = tgt.lstrip("/") if tgt.startswith("/") else os.path.normpath("xl/drawings/" + tgt).replace("\\", "/")
            img = Image.open(io.BytesIO(z.read(path)))
            sr = re.search(r"<a:srcRect[^>]*>", x)
            if sr:
                g = {k: int(v) / 100000 for k, v in re.findall(r'\s([ltrb])="(\d+)"', sr.group(0))}
                W, H = img.size
                img = img.crop((round(W * g.get("l", 0)), round(H * g.get("t", 0)),
                                round(W * (1 - g.get("r", 0))), round(H * (1 - g.get("b", 0)))))
            out = io.BytesIO()
            img.save(out, "PNG", optimize=True)
            return out.getvalue()
    except Exception:
        return None
    return None


def parse_month(data):
    """An existing month workbook (tabs d-m) -> (tabs, signature_png or None)."""
    wb = openpyxl.load_workbook(io.BytesIO(data))
    tabs = []
    for ws in wb.worksheets:
        m = re.search(r"(\d{1,2}/\d{1,2}/\d{4})", _s(ws["A1"].value))
        if not m or _s(ws["A3"].value).strip() != "Card Number":
            continue
        tc = ws.sheet_properties.tabColor
        sent = bool(tc is not None and tc.type == "rgb" and str(tc.rgb).upper().endswith("00B050"))
        cards = _read_cards(ws, 4, lambda ws, r: "เรือเอก" in _s(ws.cell(r, 8).value)
                            or not any(_s(ws.cell(r, c).value).strip() for c in range(1, 15))
                            and not any(_s(ws.cell(r + k, c).value).strip() for k in (1, 2) for c in (1, 8)))
        tabs.append({"date": _dmy(m.group(1)), "sent": sent, "cards": cards})
    return tabs, _signature_from_zip(data)


# ---------------------------------------------------------------- rules
def counts(cards):
    ships = {c["ship"].strip().upper() for c in cards if c["ship"].strip()}
    k = {"EN": 0, "SA": 0, "SH": 0}
    for c in cards:
        for e in c["ev"]:
            t = e[0].strip()[:2].upper()
            if t in k:
                k[t] += 1
    return len(ships), k["EN"], k["SA"], k["SH"]


def counts_text(cards):
    n, en, sa, sh = counts(cards)
    return f"รวมทั้งหมด {n} ลำ    เข้า {en} ครั้ง   ออก {sa} ครั้ง   เลื่อน {sh} ครั้ง"


def _ev_date(when, year):
    m = re.search(r"(\d{1,2})-(\d{1,2})-\d{1,2}:\d{2}", when)
    if not m:
        return None
    try:
        return date(year, int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def suggest_days(d0, d1, cards):
    """Default tab date for each card when one file covers several days (weekend/holiday):
    the day after the card's last movement, kept inside the file's range and never going backwards."""
    a, b = date.fromisoformat(d0), date.fromisoformat(d1)
    out, prev = [], a
    for c in cards:
        ds = [x for x in (_ev_date(e[0], a.year) for e in c["ev"]) if x]
        if ds and ds[-1].month == 12 and a.month == 1:
            ds = [x.replace(year=a.year - 1) if x.month == 12 else x for x in ds]
        g = max(ds) + timedelta(days=1) if ds else b
        g = min(max(g, a), b)
        g = max(g, prev)
        out.append(g.isoformat())
        prev = g
    return out


def _agent_size(name):
    s = name.strip()
    letters = [ch for ch in s if ch.isalpha()]
    upper = letters and all(ch.isupper() for ch in letters if ch.isascii())
    return 12 if upper or len(s) <= 30 else 14


def _agent_wraps(name, size):
    n = len(name.strip())
    return n >= 36 if size == 12 else n > 37


# ---------------------------------------------------------------- writing
def _style(dst, src):
    dst._style = copy.copy(src._style)


def _fill_sheet(ws, tpl, tab, sig_png):
    # clear everything below the header that came with the copied template
    for mr in [m for m in ws.merged_cells.ranges if m.min_row >= 4]:
        ws.unmerge_cells(str(mr))
    blank = tpl.cell(60, 20)
    for r in range(4, max(ws.max_row, 40) + 1):
        ws.row_dimensions[r].height = None
        for c in range(1, 15):
            cell = ws.cell(r, c)
            cell.value = None
            _style(cell, blank)

    d = date.fromisoformat(tab["date"])
    ws["A1"].value = f"{TITLE}\n{d.day}/{d.month}/{d.year}"
    ws["D2"].value = counts_text(tab["cards"])

    r = 4
    for card in tab["cards"]:
        ev = card["ev"] or [["", "", "", "", ""]]
        n = len(ev)
        size = _agent_size(card["agent"])
        for i, e in enumerate(ev):
            row = r + i
            proto = _FIRST if i == 0 else (_LAST if i == n - 1 else _MID)
            for c in range(1, 15):
                src = tpl.cell(_FIRST if c >= 8 else proto, c)
                _style(ws.cell(row, c), src)
            if n == 1:   # a one-row card is closed at the bottom too
                dcell = ws.cell(row, 4)
                b = copy.copy(dcell.border)
                dcell.border = Border(left=b.left, right=b.right, top=b.top, bottom=Side(style="thin"))
            h = 20.25
            if n == 1 and _agent_wraps(card["agent"], size):
                h = 34.5 if size == 12 else 40.5
            ws.row_dimensions[row].height = h
            ws.cell(row, 8).value, ws.cell(row, 9).value, ws.cell(row, 10).value = e[0], e[1], e[2]
            ws.cell(row, 13).value, ws.cell(row, 14).value = e[3], e[4]
            ws.merge_cells(start_row=row, start_column=10, end_row=row, end_column=12)
        top = ws.cell(r, 1)
        ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value = card["no"], card["agent"], card["ship"]
        ws.cell(r, 5).value, ws.cell(r, 6).value, ws.cell(r, 7).value = card["imo"], card["cs"], card["loa"]
        f = copy.copy(ws.cell(r, 2).font)
        f.sz = size
        ws.cell(r, 2).font = f
        last = r + n - 1
        ws.merge_cells(start_row=r, start_column=3, end_row=last, end_column=4)
        if n > 1:
            for c in (1, 2, 5, 6, 7):
                ws.merge_cells(start_row=r, start_column=c, end_row=last, end_column=c)
        r = last + 1
        del top
    L = r - 1

    # signature block (texts, styles, heights, merges) + picture
    for k in range(_SIG_N):
        src_r, dst_r = _SIG0 + k, L + 1 + k
        ws.row_dimensions[dst_r].height = tpl.row_dimensions[src_r].height
        for c in range(1, 15):
            s, t = tpl.cell(src_r, c), ws.cell(dst_r, c)
            _style(t, s)
            t.value = s.value
    for off, c0, c1 in _SIG_MERGES:
        ws.merge_cells(start_row=L + off, start_column=c0, end_row=L + off, end_column=c1)
    if sig_png:
        img = XLImage(io.BytesIO(sig_png))
        img.anchor = OneCellAnchor(
            _from=AnchorMarker(col=_SIG_IMG["col"], colOff=_SIG_IMG["col_off"], row=L, rowOff=_SIG_IMG["row_off"]),
            ext=XDRPositiveSize2D(_SIG_IMG["cx"], _SIG_IMG["cy"]))
        ws.add_image(img)

    ws.sheet_properties.tabColor = Color(rgb=SENT_RGB) if tab.get("sent") else Color(theme=0)


def build_month(tabs, sig_png=None):
    """tabs (any order) -> xlsx bytes, one sheet per day named d-m, oldest first."""
    wb = openpyxl.load_workbook(TEMPLATE)
    tpl = wb["TPL"]
    zoom = tpl.sheet_view.zoomScale
    for tab in sorted(tabs, key=lambda t: t["date"]):
        ws = wb.copy_worksheet(tpl)
        ws.title = tab_name(tab["date"])
        ws.sheet_view.zoomScale = zoom
        ws.sheet_view.tabSelected = False
        _fill_sheet(ws, tpl, tab, sig_png)
    wb.remove(tpl)
    if wb.worksheets:
        wb.active = len(wb.worksheets) - 1
        wb.worksheets[-1].sheet_view.tabSelected = True
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- storage (the app's Google Sheet, apps_script/Code.gs)
import base64 as _b64
import json as _json
import time as _time
import urllib.parse as _up
import urllib.request as _ur

COL_DAYS, COL_CFG = "closeDays", "reportCfg"


class SheetStore:
    """Days are docs in collection closeDays (id = ISO date); the signature picture is reportCfg/signature."""

    def __init__(self, url, key):
        self.url, self.key = url.strip(), key.strip()

    def _get(self, **q):
        q["key"] = self.key
        sep = "&" if "?" in self.url else "?"
        with _ur.urlopen(self.url + sep + _up.urlencode(q), timeout=60) as r:
            j = _json.loads(r.read().decode("utf-8"))
        if not j.get("ok"):
            raise RuntimeError("รหัสเข้าใช้ไม่ถูกต้อง" if j.get("error") == "bad_key" else str(j.get("error")))
        return j

    def _post(self, ops):
        body = _json.dumps({"key": self.key, "ops": ops}).encode("utf-8")
        req = _ur.Request(self.url, data=body, headers={"Content-Type": "text/plain;charset=utf-8"})
        with _ur.urlopen(req, timeout=60) as r:
            j = _json.loads(r.read().decode("utf-8"))
        if not j.get("ok"):
            raise RuntimeError(str(j.get("error")))
        return j

    def load(self):
        """-> (days {iso: doc}, signature_png or None)"""
        docs = self._get(action="load", since=0)["docs"]
        days, sig = {}, None
        for d in docs:
            if d.get("deleted") or not d.get("doc"):
                continue
            if d["col"] == COL_DAYS:
                days[d["id"]] = d["doc"]
            elif d["col"] == COL_CFG and d["id"] == "signature":
                try:
                    sig = _b64.b64decode(d["doc"].get("png", ""))
                except Exception:
                    sig = None
        return days, sig

    def put_days(self, days):
        now = int(_time.time() * 1000)
        ops = [{"op": "put", "col": COL_DAYS, "id": t["date"], "doc": t, "updated": now} for t in days]
        for i in range(0, len(ops), 20):
            self._post(ops[i:i + 20])

    def delete_day(self, iso):
        self._post([{"op": "del", "col": COL_DAYS, "id": iso, "updated": int(_time.time() * 1000)}])

    def put_signature(self, png):
        b = _b64.b64encode(png).decode("ascii")
        if len(b) > 48000:
            raise ValueError("รูปลายเซ็นใหญ่เกินไป (ต้องไม่เกินประมาณ 35 KB) ลองครอปให้เหลือเฉพาะลายเซ็น")
        self._post([{"op": "put", "col": COL_CFG, "id": "signature", "doc": {"png": b},
                     "updated": int(_time.time() * 1000)}])


def shrink_signature(data, max_w=600):
    """Any picture -> PNG small enough to keep in one sheet cell."""
    from PIL import Image
    img = Image.open(io.BytesIO(data))
    if img.width > max_w:
        img = img.resize((max_w, round(img.height * max_w / img.width)))
    out = io.BytesIO()
    img.save(out, "PNG", optimize=True)
    return out.getvalue()
